#!/usr/bin/env python3
"""Show what the slide builds look like, since a static preview cannot.

Usage:
  python anim_preview.py WORK_DIR [--pptx out/beautified.pptx] [--width 960] [--hold 0.8]

For every page that has builds, renders the page as it looks before the first build, after
the first, after the second ... and writes
  out/anim/storyboard.png   one row per animated page, its states left to right
  out/anim/preview.gif      the deck played through, each build one beat
Needs the same LibreOffice + pdftoppm as render.py. Renders the deck once per build depth,
so a deck whose busiest page has 4 builds costs 5 renders.
"""
import argparse
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

import render as RD  # noqa: E402


def build_groups(slide):
    """Shape ids per build, in order, read back from the slide's timing."""
    t = slide._element.find(qn("p:timing"))
    if t is None:
        return []
    seq = next(t.iter(qn("p:seq")), None)
    if seq is None:
        return []
    groups = []
    main = seq.find(qn("p:cTn")).find(qn("p:childTnLst"))
    for outer in main.findall(qn("p:par")):
        click = []
        for ctn in outer.iter(qn("p:cTn")):
            nt = ctn.get("nodeType")
            if nt not in ("clickEffect", "withEffect", "afterEffect"):
                continue
            ids = {int(s.get("spid")) for s in ctn.iter(qn("p:spTgt"))}
            delay = next(ctn.find(qn("p:stCondLst")).iter(qn("p:cond"))).get("delay", "0")
            click.append((int(delay) if delay.isdigit() else 0, ids))
        # auto pages: one click holding staggered effects; split them by start time
        by_delay = {}
        for d, ids in click:
            by_delay.setdefault(d, set()).update(ids)
        groups += [by_delay[d] for d in sorted(by_delay)]
    return groups


def state_deck(src, depth, dst):
    """Copy of the deck where each page shows only its first `depth` builds."""
    prs = Presentation(str(src))
    for slide in prs.slides:
        groups = build_groups(slide)
        hide = set().union(*groups[depth:]) if len(groups) > depth else set()
        for sh in list(slide.shapes):
            if sh.shape_id in hide:
                sh._element.getparent().remove(sh._element)
        t = slide._element.find(qn("p:timing"))
        if t is not None:
            slide._element.remove(t)
    prs.save(str(dst))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("work")
    ap.add_argument("--pptx")
    ap.add_argument("--width", type=int, default=960)
    ap.add_argument("--hold", type=float, default=0.8, help="seconds per build in the gif")
    a = ap.parse_args()
    work = Path(a.work)
    src = Path(a.pptx) if a.pptx else work / "out" / "beautified.pptx"
    out = work / "out" / "anim"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    prs = Presentation(str(src))
    counts = [len(build_groups(s)) for s in prs.slides]
    deepest = max(counts) if counts else 0
    if deepest == 0:
        print("no builds in this deck (animation mode none or transition); nothing to preview")
        return 0
    frames = {}  # depth -> [png paths]
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for depth in range(deepest + 1):
            deck = tmp / f"d{depth}.pptx"
            state_deck(src, depth, deck)
            pngs = RD.render(str(deck), str(tmp / f"r{depth}"), dpi=110, prefix="s")
            frames[depth] = [Image.open(p).convert("RGB") for p in sorted(pngs)]
    W = a.width

    def fit(im):
        return im.resize((W, round(im.height * W / im.width)), Image.LANCZOS)

    # gif: every page, with each of its builds as its own beat
    seq, durs = [], []
    for i, n in enumerate(counts):
        for depth in range(n + 1):
            seq.append(fit(frames[depth][i]))
            durs.append(int(a.hold * 1000) if depth < n else int(a.hold * 2000))
    seq[0].save(out / "preview.gif", save_all=True, append_images=seq[1:], duration=durs, loop=0,
                optimize=True)
    # storyboard: animated pages only
    tw = 360
    rows = [(i, n) for i, n in enumerate(counts) if n]
    th = round(frames[0][0].height * tw / frames[0][0].width)
    pad, lab = 10, 26
    cols = deepest + 1
    sb = Image.new("RGB", (pad + cols * (tw + pad), pad + len(rows) * (th + lab + pad)), (236, 238, 240))
    d = ImageDraw.Draw(sb)
    for r, (i, n) in enumerate(rows):
        y = pad + r * (th + lab + pad)
        for depth in range(n + 1):
            x = pad + depth * (tw + pad)
            d.text((x, y + 6), f"P{i + 1}  " + ("start" if depth == 0 else f"build {depth}/{n}"), fill=(60, 60, 60))
            sb.paste(frames[depth][i].resize((tw, th), Image.LANCZOS), (x, y + lab))
    sb.save(out / "storyboard.png")
    print(f"wrote {out / 'storyboard.png'} and {out / 'preview.gif'} ({sum(counts)} builds on {len(rows)} pages)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
