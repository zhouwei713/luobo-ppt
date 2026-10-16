#!/usr/bin/env python3
"""Step 1: read the source deck.

Usage:
  python extract_deck.py source.pptx WORK_DIR [--no-render]

Writes into WORK_DIR:
  source.pptx          normalized copy (fixes zips LibreOffice cannot open)
  raw.json             every shape: position (inches), fill, text runs, pictures, notes
  source_text.json     every text unit per slide (used by qa_check for fidelity)
  deck.draft.json      first guess of the semantic deck; Claude edits it into deck.json
  media/               pictures from the deck (svg kept as svg, text inside svg listed)
  source_png/          one PNG per slide + sheet.png (unless --no-render)
"""
import argparse
import re
import sys
import html
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import normalize_pptx, save_json  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu  # noqa: E402

EMU_IN = 914400
NOTES_RE = re.compile(r"notes=\{`(.*?)`\}", re.S)


def inch(v):
    return round((v or 0) / EMU_IN, 3)


def fill_hex(shape):
    try:
        if shape.fill.type == 1:  # solid
            return str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def run_info(run):
    f = run.font
    color = None
    try:
        if f.color and f.color.type is not None and f.color.rgb is not None:
            color = str(f.color.rgb)
    except Exception:
        pass
    return {"text": run.text, "bold": bool(f.bold), "size": f.size.pt if f.size else None,
            "color": color}


def para_text_marked(p):
    """Paragraph text with **bold** markers for partially-bold runs.

    Soft line breaks (<a:br>, Shift+Enter) come out as "\v", the same convention
    python-pptx uses, so they stay distinct from the "\n" that joins paragraphs.
    """
    runs = [r for r in p.runs if r.text]
    if not runs:
        return p.text if hasattr(p, "text") else ""
    all_bold = all(r.font.bold for r in runs)
    marked = {r._r: f"**{r.text}**" if not all_bold and r.font.bold and r.text.strip() else r.text
              for r in runs}
    out = "".join("\v" if child.tag == qn("a:br") else marked.get(child, "") for child in p._p)
    return out.replace("****", "")


def svg_texts(data):
    try:
        s = data.decode("utf-8", "ignore")
    except Exception:
        return []
    return [html.unescape(re.sub(r"<[^>]+>", "", m)).strip()
            for m in re.findall(r"<text[^>]*>(.*?)</text>", s, re.S) if m.strip()]


def walk(shapes, out, media_dir, slide_n, counter):
    for sh in shapes:
        counter[0] += 1
        base = {"id": f"s{slide_n}_{counter[0]}", "name": sh.name,
                "x": inch(sh.left), "y": inch(sh.top), "w": inch(sh.width), "h": inch(sh.height)}
        st = sh.shape_type
        if st == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, out, media_dir, slide_n, counter)
            continue
        if st == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
            try:
                img = sh.image
                ext = img.ext
                blob = img.blob
            except Exception:
                img, ext, blob = None, "bin", b""
            # SVG pictures: python-pptx exposes the PNG fallback; look for svgBlip
            svg_blob = None
            try:
                from lxml import etree
                xml = etree.tostring(sh._element).decode()
                m = re.search(r'svgBlip[^>]*r:embed="([^"]+)"', xml) or \
                    re.search(r'blip[^>]*r:embed="([^"]+)"', xml)
                if m:
                    part = sh.part.related_part(m.group(1))
                    if part.partname.endswith(".svg"):
                        svg_blob = part.blob
            except Exception:
                pass
            if svg_blob is not None:
                ext, blob = "svg", svg_blob
            path = media_dir / f"{base['id']}.{ext}"
            if blob:
                path.write_bytes(blob)
            item = dict(base, kind="picture", path=str(path.name))
            if ext == "svg":
                item["svg_texts"] = svg_texts(blob)
                # a PNG copy the layouts can place (python-pptx cannot insert SVG)
                from common import svg_to_png
                png = media_dir / f"{base['id']}.png"
                if svg_to_png(path, png):
                    item["png"] = png.name
            out.append(item)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = [[c.text for c in r.cells] for r in sh.table.rows]
            out.append(dict(base, kind="table", rows=rows))
            continue
        item = dict(base, kind="shape", fill=fill_hex(sh))
        if getattr(sh, "is_placeholder", False):
            try:
                item["placeholder"] = str(sh.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
            except Exception:
                item["placeholder"] = "body"
        try:
            item["auto_shape"] = str(sh.auto_shape_type).split(".")[-1].split(" ")[0]
        except Exception:
            pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            paras = []
            for p in sh.text_frame.paragraphs:
                if not p.text.strip():
                    continue
                paras.append({"text": para_text_marked(p), "level": p.level,
                              "runs": [run_info(r) for r in p.runs if r.text]})
            sizes = [r["size"] for p in paras for r in p["runs"] if r["size"]]
            item.update(kind="text", text="\n".join(p["text"] for p in paras), paragraphs=paras,
                        size=max(sizes) if sizes else None,
                        bold=all(r["bold"] for p in paras for r in p["runs"]) if paras else False,
                        color=next((r["color"] for p in paras for r in p["runs"] if r["color"]), None))
        out.append(item)


def slide_notes(slide):
    txt = ""
    if slide.has_notes_slide:
        txt = slide.notes_slide.notes_text_frame.text.strip()
    if not txt:
        # some generators hide notes inside a descr attribute: notes={`...`}
        from lxml import etree
        xml = etree.tostring(slide._element).decode()
        m = NOTES_RE.search(html.unescape(xml))
        if m:
            txt = m.group(1).strip()
    return txt


def container_of(t, shapes, sw, sh_):
    best = None
    for s in shapes:
        if s["kind"] != "shape" or not s.get("fill") or s is t:
            continue
        if s["w"] >= sw * 0.95 and s["h"] >= sh_ * 0.95:
            continue  # slide background
        if s["x"] - 0.05 <= t["x"] and s["y"] - 0.05 <= t["y"] and \
           t["x"] + t["w"] <= s["x"] + s["w"] + 0.05 and t["y"] + t["h"] <= s["y"] + s["h"] + 0.05:
            if best is None or s["w"] * s["h"] < best["w"] * best["h"]:
                best = s
    return best


def guess_role(t, sw, sh_, max_size_top):
    txt = t["text"].strip()
    ph = t.get("placeholder")
    if ph in ("title", "center_title"):
        return "title"
    if ph == "subtitle":
        return "subtitle"
    size = t.get("size") or 14
    if re.fullmatch(r"\d{1,3}\s*/\s*\d{1,3}", txt):
        return "page"
    if t["y"] > sh_ * 0.82 and size <= 12:
        return "footer"
    if t["y"] < sh_ * 0.25 and size >= max_size_top - 0.1 and size >= 24:
        return "title"
    if len(txt) <= 6 and re.search(r"\d", txt) and size >= 40:
        return "number"
    if txt in ("「", "」", "“", "”", "\""):
        return "decor"
    return "text"


def guess_type(n, total, slide_bg, units):
    roles = [u["role"] for u in units]
    nums = roles.count("number")
    if n == 1:
        return "cover"
    if n == total and slide_bg and slide_bg not in ("FFFFFF",):
        return "closing"
    if nums >= 2:
        return "stats"
    if nums == 1:
        return "big_number"
    return "content"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("work")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    work = Path(a.work)
    work.mkdir(parents=True, exist_ok=True)
    media = work / "media"
    media.mkdir(exist_ok=True)
    src = work / "source.pptx"
    normalize_pptx(a.pptx, src)
    # equations typed with the equation editor become plain formula text, so they are
    # extracted, previewed and checked like every other line (LibreOffice draws none of them)
    from omml import flatten_math
    eqs = flatten_math(src)
    if eqs:
        n_eq = sum(len(v) for v in eqs.values())
        print(f"NOTE 原稿里有 {n_eq} 个公式对象，已转成普通文字公式（任何软件都能显示，也能直接改字）：")
        for k in sorted(eqs):
            print(f"  slide{k}.xml: " + " ｜ ".join(eqs[k]))
    prs = Presentation(str(src))
    sw, sh_ = inch(prs.slide_width), inch(prs.slide_height)
    raw = {"slide_w": sw, "slide_h": sh_, "slides": []}
    source_text = {"slides": []}
    draft = {"meta": {"title": "", "series": "", "footer": "", "total": len(prs.slides),
                      "topic": "", "anchor": {"description": "", "glyph": ""}},
             "slides": []}
    total = len(prs.slides)
    for n, slide in enumerate(prs.slides, 1):
        shapes = []
        walk(slide.shapes, shapes, media, n, [0])
        bg = None
        for s in shapes:
            if s["kind"] == "shape" and s.get("fill") and s["w"] >= sw * 0.95 and s["h"] >= sh_ * 0.95:
                bg = s["fill"]
        notes = slide_notes(slide)
        raw["slides"].append({"n": n, "bg": bg, "notes": notes, "shapes": shapes})

        texts = [s for s in shapes if s["kind"] == "text"]
        units_src = []
        for t in texts:
            # one unit per paragraph so layouts may split a bullet list into cards
            for k, para in enumerate(t["paragraphs"]):
                units_src.append({"id": f"{t['id']}_p{k}", "text": para["text"]})
        for p in shapes:
            if p["kind"] == "picture":
                for i, st in enumerate(p.get("svg_texts", [])):
                    units_src.append({"id": f"{p['id']}_svg{i}", "text": st, "in_image": True})
            if p["kind"] == "table":
                for r in p["rows"]:
                    for c in r:
                        if c.strip():
                            units_src.append({"id": p["id"] + "_cell", "text": c})
        source_text["slides"].append({"n": n, "units": units_src})

        top_sizes = [t.get("size") or 0 for t in texts if t["y"] < sh_ * 0.25]
        max_top = max(top_sizes) if top_sizes else 0
        units = []
        for t in sorted(texts, key=lambda s: (round(s["y"], 1), s["x"])):
            c = container_of(t, shapes, sw, sh_)
            units.append({"id": t["id"], "text": t["text"],
                          "paragraphs": [p_["text"] for p_ in t["paragraphs"]] if len(t["paragraphs"]) > 1 else None,
                          "role": guess_role(t, sw, sh_, max_top),
                          "size": t.get("size"), "bold": t.get("bold"), "color": t.get("color"),
                          "x": t["x"], "y": t["y"], "w": t["w"],
                          "container": {"id": c["id"], "fill": c["fill"]} if c else None})
        pics = [{"id": p["id"], "path": p["path"], "svg_texts": p.get("svg_texts", []),
                 **({"png": p["png"]} if p.get("png") else {}),
                 "x": p["x"], "y": p["y"], "w": p["w"], "h": p["h"]}
                for p in shapes if p["kind"] == "picture"]
        title = next((u["text"] for u in units if u["role"] == "title"), "")
        draft["slides"].append({"n": n, "type": guess_type(n, total, bg, units), "title": title,
                                "notes": notes, "units": units, "pictures": pics})
        for u in units:
            if u["role"] == "footer" and not draft["meta"]["footer"]:
                draft["meta"]["footer"] = u["text"]

    save_json(raw, work / "raw.json")
    save_json(source_text, work / "source_text.json")
    save_json(draft, work / "deck.draft.json")
    print(f"slides: {total}  size: {sw}x{sh_} in")
    print(f"wrote {work/'raw.json'}, {work/'source_text.json'}, {work/'deck.draft.json'}")
    if not a.no_render:
        try:
            from render import render
            pngs = render(src, work / "source_png", dpi=80, sheet=work / "source_png" / "sheet.png")
            print(f"rendered {len(pngs)} slides -> {work/'source_png'}/sheet.png")
        except SystemExit as e:
            print(f"NOTE 原稿预览图未生成：{e}")


if __name__ == "__main__":
    main()
