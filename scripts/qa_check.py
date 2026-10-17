#!/usr/bin/env python3
"""Step 6: automatic checks on the beautified deck.

Usage:
  python qa_check.py WORK_DIR [--pptx WORK_DIR/out/beautified.pptx]

Checks
  fidelity   every source text unit still appears on the same slide (whitespace and
             **bold** marks ignored). Missing text = ERROR. Text found only inside a
             source picture (svg) = WARN if missing. Added text (header, page no.) = INFO.
  overflow   text needs more height than its box (measured with CJK font metrics) = ERROR
  overlap    two text boxes intersect = ERROR; text over the opaque part of an
             illustration = WARN
  bounds     element outside the slide = WARN
  contrast   text vs what is under it below WCAG thresholds = WARN
  tiny       font below 10pt = WARN
  animation  a build pointing at a missing shape = ERROR; more than 6 clicks on one page = WARN
  chart      every chart of the original (or of deck.json) is on the slide with the same categories,
             series names and values; numbers read the same (35% stays 35%, not 35); chart and
             axis titles appear on the chart or the slide = ERROR otherwise
  fine-print notes, sources, unit lines and footnote marks (*, ※, ¹) kept exactly, marks included
             = ERROR otherwise. Both are listed side by side in out/fine_print.md
Writes WORK_DIR/out/qa_report.json. Exit code 1 when any ERROR.
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (load_json, save_json, load_style, wrap_lines, contrast, rgb_to_hex, text_width_pt, DISPLAY_LATIN, WIDTH_SAFETY, LINE_PITCH,  # noqa: E402
                    mix, CANVAS_W, CANVAS_H)

from pptx import Presentation  # noqa: E402


def norm(s):
    # Ignore list markers, but require space after a dash to preserve negative numbers.
    s = re.sub(r"(?m)^[^\S\r\n]*(?:[·•●▪◦*][^\S\r\n]*|[\-–][^\S\r\n]+)", "", (s or "").replace("**", ""))
    s = re.sub(r"\s+", "", s)
    return s.replace("“", "\"").replace("”", "\"")


def slide_texts(prs):
    out = []
    for sl in prs.slides:
        units = []
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                units.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            units.append(cell.text)
        out.append(units)
    return out


def rect_inter(a, b):
    x0, y0 = max(a[0], b[0]), max(a[1], b[1])
    x1, y1 = min(a[2], b[2]), min(a[3], b[3])
    return max(0, x1 - x0) * max(0, y1 - y0)


_ALPHA_BOX = {}


def opaque_box(path, placed):
    """Map the non-transparent bbox of an image into slide inches."""
    from PIL import Image
    if path not in _ALPHA_BOX:
        try:
            im = Image.open(path)
            if im.mode != "RGBA":
                _ALPHA_BOX[path] = None
            else:
                bb = im.getchannel("A").point(lambda v: 255 if v > 40 else 0).getbbox()
                _ALPHA_BOX[path] = (bb, im.size) if bb else None
        except Exception:
            _ALPHA_BOX[path] = None
    info = _ALPHA_BOX[path]
    x, y, w, h = placed
    if not info:
        return (x, y, x + w, y + h)
    (bx0, by0, bx1, by1), (iw, ih) = info
    return (x + w * bx0 / iw, y + h * by0 / ih, x + w * bx1 / iw, y + h * by1 / ih)


_BG_CACHE = {}


def sample_bg(path, box):
    from PIL import Image
    if path not in _BG_CACHE:
        _BG_CACHE[path] = Image.open(path).convert("RGB").resize((400, 225))
    im = _BG_CACHE[path]
    x0, y0, x1, y1 = box
    sx, sy = 400 / CANVAS_W, 225 / CANVAS_H
    crop = im.crop((int(max(0, x0) * sx), int(max(0, y0) * sy), int(min(CANVAS_W, x1) * sx) + 1, int(min(CANVAS_H, y1) * sy) + 1))
    import numpy as np
    arr = np.asarray(crop).reshape(-1, 3)
    if not len(arr):
        return "FFFFFF"
    return rgb_to_hex(arr.mean(0))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("work")
    ap.add_argument("--pptx")
    a = ap.parse_args()
    work = Path(a.work)
    pptx = Path(a.pptx) if a.pptx else work / "out" / "beautified.pptx"
    style = load_style(work / "style.json")
    dump = load_json(pptx.parent / "layout_dump.json")
    deck = load_json(work / "deck.json")
    src_path = work / "source_text.json"
    gen_mode = not src_path.exists()
    if gen_mode:
        # generated deck: the confirmed deck.json is the reference text
        import sys as _s
        _s.path.insert(0, str(Path(__file__).resolve().parent))
        from outline_check import screen_text
        src = {"slides": [{"n": sl["n"], "units": [{"id": f"d{sl['n']}_{i}", "text": t}
                                                   for i, t in enumerate(screen_text(sl)) if t.strip()]}
                          for sl in deck["slides"]]}
    else:
        src = load_json(src_path)
    prs = Presentation(str(pptx))
    out_units = slide_texts(prs)
    issues = []

    def add(level, n, kind, msg):
        issues.append({"level": level, "slide": n, "check": kind, "msg": msg})

    # ---------------- fidelity
    pic_stems = {sd["n"]: {Path(e["path"]).stem for e in sd["elements"] if e["t"] == "image" and e.get("path")}
                 for sd in dump["slides"]}
    if len(prs.slides) != len(src["slides"]):
        add("ERROR", 0, "fidelity", f"slide count {len(prs.slides)} != source {len(src['slides'])}")
    for i, ss in enumerate(src["slides"]):
        n = ss["n"]
        if i >= len(out_units):
            break
        blob = norm("".join(out_units[i]))
        src_norms = []
        for u in ss["units"]:
            t = norm(u["text"])
            src_norms.append(t)
            if not t:
                continue
            if len(t) == 1 and not re.match(r"[\w一-鿿]", t):
                continue  # decorative punctuation such as 「 replaced by an icon
            if t not in blob and u.get("in_image") and str(u.get("id", "")).split("_svg")[0] in pic_stems.get(n, set()):
                continue  # the text lives inside the original picture, which is on the slide
            if t not in blob:
                lvl = "WARN" if u.get("in_image") else "ERROR"
                what = "deck.json text not on slide" if gen_mode else "missing text"
                add(lvl, n, "fidelity", f"{what}: {u['text'][:60]!r}")
        allsrc = "".join(src_norms)
        for t in out_units[i]:
            tn = norm(t)
            if tn and tn not in allsrc:
                add("INFO", n, "added", f"added text: {t[:40]!r}")

    # ---------------- charts and fine print: the body text can be right while a unit, an axis
    # title, a series name or a footnote mark has changed the meaning. Checked on their own,
    # and listed side by side in out/fine_print.md for a last look before delivery.
    import chartcheck as CC
    fp = ["# 图表和小字对照", "",
          "正文逐字核对之外，单位、坐标轴、系列名、注释和来源单独对一遍。左边是原稿（生成时是 deck.json），右边是成品。", ""]
    for i, ss in enumerate(src["slides"]):
        if i >= len(prs.slides):
            break
        n = ss["n"]
        slide = prs.slides[i]
        text_blob = "".join(out_units[i])
        if gen_mode:
            dsl = deck["slides"][i] if i < len(deck["slides"]) else {}
            refs = [CC.from_spec(dsl["chart"])] if dsl.get("chart") else []
            units_all = ss["units"]
        else:
            refs = [c for c in ss.get("charts", []) if c]
            units_all = ss["units"]
        outs = CC.charts_on_slide(slide)
        section = []
        if len(outs) < len(refs):
            add("ERROR", n, "chart", f"原稿有 {len(refs)} 个图表，成品只有 {len(outs)} 个；图表要写成 chart 页型，"
                                     "数据从 deck.draft.json 的 charts 里照抄")
        left = list(outs)
        for k, ref in enumerate(refs):
            if not left:
                break
            best = max(left, key=lambda o: len(set(o["categories"]) & set(ref["categories"])))
            left.remove(best)
            probs, rows = CC.compare(ref, best, text_blob)
            for lvl, msg in probs:
                add(lvl, n, "chart", msg)
            section.append(f"图表 {k + 1}" + ("，有问题" if probs else "，一致"))
            section.append("")
            section.append("| 项目 | 原稿 | 成品 |")
            section.append("|---|---|---|")
            section += [f"| {a} | {b} | {c} |" for a, b, c in rows]
            section.append("")
        smalls = CC.fine_print_units(units_all)
        sblob = CC.strict(text_blob)
        for u in smalls:
            ok = CC.strict(u["text"]) in sblob
            if not ok and norm(u["text"]) in norm(text_blob):
                add("ERROR", n, "fine-print", f"小字被改了（注释符号、星号或括号没带过来）：{u['text'][:50]!r}")
        if smalls:
            section.append("| 小字 | 成品里 |")
            section.append("|---|---|")
            for u in smalls:
                ok = CC.strict(u["text"]) in sblob
                section.append(f"| {u['text'][:80].replace('|', '/')} | {'原样' if ok else '变了或缺了'} |")
            section.append("")
        if section:
            fp += [f"## 第 {n} 页", ""] + section
    if len(fp) == 4:
        fp.append("这份稿子没有图表，也没有注释、来源、单位这类小字。")
    (pptx.parent / "fine_print.md").write_text("\n".join(fp) + "\n", encoding="utf-8")

    # ---------------- geometry
    for sd in dump["slides"]:
        n = sd["n"]
        els = sd["elements"]
        texts = [e for e in els if e["t"] == "text" and e.get("role") != "decor"]  # watermarks are decoration
        for e in els:
            if e.get("role") == "background":
                continue
            x, y = e["x"], e["y"]
            w, h = e.get("w", e.get("size", 0)), e.get("h", e.get("size", 0))
            if x < -0.01 or y < -0.01 or x + w > CANVAS_W + 0.01 or y + h > CANVAS_H + 0.01:
                add("WARN", n, "bounds", f"{e['t']} outside slide at ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
        for e in texts:
            ls_ = DISPLAY_LATIN if e.get("font") == "display" else 1.0
            lines = wrap_lines(e["text"], e["size"], e["w"] * 72 * (WIDTH_SAFETY + 0.02), e.get("bold", False), ls_)
            need = len(lines) * e["size"] * e.get("ls", 1.2) * LINE_PITCH / 72
            if need > e["h"] * 1.08 + 0.02:
                add("ERROR", n, "overflow", f"{e['text'][:30]!r} needs {need:.2f}in, box {e['h']:.2f}in "
                                            f"({len(lines)} lines at {e['size']}pt)")
            if e["size"] < 10:
                add("WARN", n, "tiny", f"{e['text'][:30]!r} at {e['size']}pt")
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a_, b_ = texts[i], texts[j]
                ra = (a_["x"], a_["y"], a_["x"] + a_["w"], a_["y"] + a_["h"])
                rb = (b_["x"], b_["y"], b_["x"] + b_["w"], b_["y"] + b_["h"])
                inter = rect_inter(ra, rb)
                small = min(a_["w"] * a_["h"], b_["w"] * b_["h"])
                if small and inter / small > 0.12:
                    add("ERROR", n, "overlap", f"text {a_['text'][:18]!r} overlaps {b_['text'][:18]!r}")
        heroes = [e for e in els if e["t"] == "image" and e.get("role") not in ("background", "decor")
                  and e.get("fit", "contain") == "contain"]
        for himg in heroes:
            placed = (himg.get("rx", himg["x"]), himg.get("ry", himg["y"]), himg.get("rw", himg["w"]), himg.get("rh", himg["h"]))
            ob = opaque_box(himg["path"], placed)
            for t in texts:
                # use where glyphs actually are: shrink the box to the measured line width
                lines = wrap_lines(t["text"], t["size"], t["w"] * 72 * 0.98, t.get("bold", False))
                lw = min(t["w"], max(text_width_pt(l, t["size"], t.get("bold", False)) for l in lines) / 72)
                x0 = t["x"] if t.get("align", "l") == "l" else (t["x"] + t["w"] - lw if t.get("align") == "r" else t["x"] + (t["w"] - lw) / 2)
                rt = (x0, t["y"], x0 + lw, t["y"] + t["h"])
                inter = rect_inter(rt, ob)
                if inter / max(0.01, lw * t["h"]) > 0.15:
                    add("WARN", n, "overlap", f"text {t['text'][:18]!r} sits on illustration {Path(himg['path']).name}")
        # contrast
        bg_img = next((e["path"] for e in els if e.get("role") == "background" and e["t"] == "image"), None)
        bg_fill = next((style.c(e["fill"]) for e in els if e.get("role") == "background" and e["t"] == "rect"), "FFFFFF")
        for idx, t in enumerate(els):
            if t["t"] != "text" or t.get("role") == "decor":
                continue
            cx, cy = t["x"] + t["w"] / 2, t["y"] + t["h"] / 2
            under = None
            for e in reversed(els[:idx]):
                if e["t"] in ("rect", "ellipse") and e.get("fill") and e.get("role") != "background":
                    if e["x"] <= cx <= e["x"] + e["w"] and e["y"] <= cy <= e["y"] + e["h"]:
                        under = e
                        break
            if under:
                behind = sample_bg(bg_img, (cx - 0.3, cy - 0.1, cx + 0.3, cy + 0.1)) if bg_img else bg_fill
                col = mix(behind, style.c(under["fill"]), under.get("alpha", 1.0))
            else:
                col = sample_bg(bg_img, (t["x"], t["y"], t["x"] + t["w"], t["y"] + t["h"])) if bg_img else bg_fill
            cr = contrast(style.c(t["color"]), col)
            large = t["size"] >= 18 or (t.get("bold") and t["size"] >= 14)
            need = 3.0 if large else 4.5
            if cr < need:
                add("WARN", n, "contrast", f"{t['text'][:24]!r} contrast {cr:.1f} < {need} (text #{style.c(t['color'])} on #{col})")

    # ---------------- rail: content stays above the takeaway pill / footer row
    for sd in dump["slides"]:
        if sd.get("master", "content") != "content":
            continue
        rail = sd.get("rail", 6.8)
        for e in sd["elements"]:
            if e.get("role") in ("background", "decor", "chrome") or e["t"] == "line":
                continue
            bottom = e["y"] + e.get("h", e.get("size", 0))
            if bottom > rail + 0.03:
                if e["t"] == "text":
                    add("ERROR", sd["n"], "rail", f"{e['text'][:24]!r} 底部 {bottom:.2f}in 越过底线 {rail:.2f}in")
                elif e["t"] != "image":
                    add("WARN", sd["n"], "rail", f"{e['t']} 底部 {bottom:.2f}in 越过底线 {rail:.2f}in")

    # ---------------- accent discipline: few, small accent moments per slide
    acc_cap = style.data.get("accent_max", {"uses": 3, "area": 0.08})
    for sd in dump["slides"]:
        uses, area = 0, 0.0
        for e in sd["elements"]:
            if e.get("role") in ("background", "decor", "chrome"):
                continue
            tok = str(e.get("fill") or e.get("color") or "")
            if e["t"] == "icon":
                tok = str(e.get("color") or "")
            if tok.startswith("accent"):
                uses += 1
                if e["t"] in ("rect", "ellipse") and e.get("fill"):
                    area += e["w"] * e["h"] * e.get("alpha", 1.0)
        share = area / (CANVAS_W * CANVAS_H)
        if uses > acc_cap["uses"] or share > acc_cap["area"]:
            add("WARN", sd["n"], "accent", f"强调色出现 {uses} 处、面积 {share:.0%}，上限 {acc_cap['uses']} 处、{acc_cap['area']:.0%}")

    # ---------------- style no-go list (style.json "forbid")
    fb = style.data.get("forbid", {})
    P = style.palette
    if fb.get("pure_white") and ("FFFFFF" in (P["bg_top"].upper(), P["bg_bottom"].upper())):
        add("WARN", 0, "forbid", "这套风格不用纯白底，bg_top / bg_bottom 换成带一点暖或冷调的浅色")
    if fb.get("pure_black") and "000000" in (P["text"].upper(), P["bg_bottom"].upper(), P["primary_dark"].upper()):
        add("WARN", 0, "forbid", "这套风格不用纯黑，换成带色相的近黑")
    for sd in dump["slides"]:
        for e in sd["elements"]:
            if e.get("role") in ("background", "decor"):
                continue
            if fb.get("shadow") and e.get("shadow") and style.depth not in ("flat", "ring"):
                add("WARN", sd["n"], "forbid", f"{e['t']} 带阴影，这套风格不用阴影")
                break
            ro = fb.get("radius_over")
            if ro is not None and e["t"] == "rect" and (e.get("radius") or 0) > ro + 0.001 and e["h"] > 0.2:
                add("WARN", sd["n"], "forbid", f"圆角 {e['radius']:.2f} 超过这套风格的上限 {ro}")
                break

    # content checks (both modes): body text size and deck-level structure
    proj = bool(deck.get("meta", {}).get("projection"))
    floor = 14 if proj else 12
    for sd in dump["slides"]:
        for e in sd["elements"]:
            if e["t"] == "text" and e["size"] < floor and len(e["text"]) > 12 and (proj or e.get("color") not in ("muted", "subtle")):
                add("WARN", sd["n"], "small-text", f"{e['text'][:24]!r} 只有 {e['size']}pt，"
                    + ("投影模式下文字不小于 14pt" if proj else "正文建议不小于 12pt"))
    if len(deck["slides"]) > 15 and not any(sl.get("type") == "section" for sl in deck["slides"]):
        add("WARN", 0, "structure", "超过 15 页，没有章节分隔页")
    # ---------------- animation: every build must point at a shape on its own page
    import animate as _A
    for i, slide in enumerate(prs.slides):
        n = dump["slides"][i]["n"] if i < len(dump["slides"]) else i + 1
        missing = _A.animated_ids(slide) - {sh.shape_id for sh in slide.shapes}
        if missing:
            add("ERROR", n, "animation", f"动画指向不存在的形状 {sorted(missing)}，重新运行 build_deck.py")
        an = dump["slides"][i].get("animation", {}) if i < len(dump["slides"]) else {}
        if an.get("mode") == "click" and an.get("builds", 0) > 6:
            add("WARN", n, "animation", f"这一页要点 {an['builds']} 次才出完，讲的时候容易乱，考虑拆页或改 auto")
    ghost = [f"{sl['n']}. {(sl.get('title') or sl.get('headline') or sl.get('question') or sl.get('quote') or '').replace('**', '')}" for sl in deck["slides"]]

    order = {"ERROR": 0, "WARN": 1, "INFO": 2}
    issues.sort(key=lambda i: (order[i["level"]], i["slide"]))
    counts = {k: sum(1 for i in issues if i["level"] == k) for k in order}
    save_json({"mode": "generate" if gen_mode else "beautify", "counts": counts, "ghost_deck": ghost,
               "issues": issues}, pptx.parent / "qa_report.json")
    for i in issues:
        if i["level"] != "INFO":
            print(f"[{i['level']}] slide {i['slide']:>2} {i['check']}: {i['msg']}")
    info = [i for i in issues if i["level"] == "INFO"]
    if info:
        print(f"[INFO] {len(info)} added text items (header, page numbers...) listed in qa_report.json")
    print(f"ERROR {counts['ERROR']}  WARN {counts['WARN']}  INFO {counts['INFO']}")
    sys.exit(1 if counts["ERROR"] else 0)


if __name__ == "__main__":
    main()
