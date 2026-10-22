#!/usr/bin/env python3
"""Build decks on a company's own PowerPoint template (.pptx / .potx).

The template's slide master stays in the output: its background, logo, bars and footer text show
on every page exactly as the company designed them. The template's theme gives the palette and the
fonts, and each layout's free area (what the master's logo, bars and footer leave over) is where
the content goes.

Usage:
  python template_kit.py read TEMPLATE WORK    # copy + analyse the template, write template.json,
                                               # style.json and template/template_card.png
  python template_kit.py apply WORK            # rebuild style.json after hand-editing template.json
  python template_kit.py card WORK             # redraw the confirmation card

template.json is meant to be read and corrected: the primary and accent colours, which layout is
used for the cover, section, closing and content pages, and each layout's free area (inches).
"""
import argparse
import colorsys
import io
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import CANVAS_H, CANVAS_W, contrast, hex_clean, load_json, mix, rel_lum, save_json  # noqa: E402

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU = 914400
TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PRESENTATION_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
# theme accents PowerPoint ships with: a deck still on these has probably never been branded
OFFICE_ACCENTS = {"4472C4", "156082", "0F9ED5", "4F81BD", "5B9BD5", "1CADE4", "E48312", "ED7D31"}
NAMES = {
    "cover": ("封面", "标题幻灯片", "title slide", "cover", "首页"),
    "section": ("节标题", "章节", "过渡", "section", "divider", "chapter"),
    "closing": ("结束", "封底", "谢谢", "致谢", "结尾", "end", "thank", "closing", "back cover"),
    "content": ("空白", "blank", "仅标题", "title only", "内容", "content"),
}
GAP = 0.15  # inches kept between template graphics and content


def _q(tag, ns=A):
    return f"{{{ns}}}{tag}"


# ------------------------------------------------------------------ opening
def open_template(src):
    """Presentation for a .pptx or .potx file (a .potx just carries another content type)."""
    data = Path(src).read_bytes()
    try:
        zin = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise SystemExit(f"{src} 不是 .pptx / .potx 文件。旧版 .ppt 或 .pot 请先在 PowerPoint 里另存为 .pptx")
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    if TEMPLATE_CT in ct:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                blob = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    blob = ct.replace(TEMPLATE_CT, PRESENTATION_CT).encode("utf-8")
                zout.writestr(item, blob)
        data = buf.getvalue()
    return Presentation(io.BytesIO(data))


def drop_slides(prs):
    """Templates often ship sample pages; the deck starts empty and keeps only masters and layouts."""
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


# ------------------------------------------------------------------ theme and colours
def theme_of(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    colors = {}
    cs = root.find(f".//{_q('clrScheme')}")
    if cs is not None:
        for el in cs:
            name = etree.QName(el).localname
            for c in el:
                tag = etree.QName(c).localname
                if tag == "srgbClr":
                    colors[name] = c.get("val").upper()
                elif tag == "sysClr":
                    colors[name] = (c.get("lastClr") or ("000000" if c.get("val") == "windowText" else "FFFFFF")).upper()
    fonts = {}
    fs = root.find(f".//{_q('fontScheme')}")
    if fs is not None:
        for kind in ("majorFont", "minorFont"):
            f = fs.find(_q(kind))
            if f is None:
                continue
            for script in ("latin", "ea"):
                el = f.find(_q(script))
                face = el.get("typeface") if el is not None else ""
                if not face and script == "ea":
                    # an empty East Asian slot names the Chinese face per script instead; Office's
                    # own default there is SimSun, which nobody picks on purpose for slides
                    for fe in f.findall(_q("font")):
                        if fe.get("script") == "Hans" and fe.get("typeface") not in ("宋体", "SimSun"):
                            face = fe.get("typeface")
                fonts[f"{kind}_{script}"] = face or ""
    return colors, fonts


def clr_map(master_el, layout_el=None):
    cm = master_el.find(_q("clrMap", P))
    out = dict(cm.attrib) if cm is not None else {}
    if layout_el is not None:
        ov = layout_el.find(f"{_q('clrMapOvr', P)}/{_q('overrideClrMapping')}")
        if ov is not None:
            out.update(ov.attrib)
    return out


def _apply_mods(hexc, el):
    r, g, b = (int(hexc[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        tag = etree.QName(m).localname
        v = int(m.get("val", "100000")) / 100000
        if tag == "lumMod":
            l *= v
        elif tag == "lumOff":
            l += v
        elif tag == "tint":  # toward white
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c * v + (1 - v) for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif tag == "shade":  # toward black
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif tag == "satMod":
            s *= v
    l, s = min(max(l, 0), 1), min(max(s, 0), 1)
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b))


def color_in(el, theme, cmap):
    """The colour a fill element (solidFill, bgRef, a gradient stop...) names, or None."""
    if el is None:
        return None
    for c in el:
        tag = etree.QName(c).localname
        base = None
        if tag == "srgbClr":
            base = c.get("val")
        elif tag == "schemeClr":
            v = c.get("val")
            v = cmap.get(v, v)
            base = theme.get(v)
        elif tag == "sysClr":
            base = c.get("lastClr")
        if base:
            return _apply_mods(base.upper(), c)
    return None


def background(el, theme, cmap):
    """('color', hex) / ('image', None) / None for a master's or layout's own background."""
    bg = el.find(f"{_q('cSld', P)}/{_q('bg', P)}")
    if bg is None:
        return None
    pr = bg.find(_q("bgPr", P))
    if pr is not None:
        if pr.find(_q("solidFill")) is not None:
            return "color", color_in(pr.find(_q("solidFill")), theme, cmap)
        gr = pr.find(_q("gradFill"))
        if gr is not None:
            stops = [color_in(gs, theme, cmap) for gs in gr.iter(_q("gs"))]
            stops = [s for s in stops if s]
            if stops:
                c = stops[0]
                for s in stops[1:]:
                    c = mix(c, s, 1 / len(stops))
                return "color", c
        if pr.find(_q("blipFill")) is not None or pr.find(_q("pattFill")) is not None:
            return "image", None
    ref = bg.find(_q("bgRef", P))
    if ref is not None:
        c = color_in(ref, theme, cmap)
        if c:
            return "color", c
    return None


def layout_bg(prs, layout, theme):
    cmap = clr_map(prs.slide_master._element, layout._element)
    got = background(layout._element, theme, cmap) or background(prs.slide_master._element, theme, cmap)
    if got and got[0] == "color" and got[1]:
        return got[1], "xml"
    if got and got[0] == "image":
        return None, "image"
    # nothing set anywhere: the page is the theme's bg1
    return theme.get(cmap.get("bg1", "lt1"), "FFFFFF"), "theme"


# ------------------------------------------------------------------ graphics the template puts on every page
def _fill_of(shape, theme, cmap):
    sppr = shape._element.find(f".//{_q('spPr', P)}")
    if sppr is not None:
        sf = sppr.find(_q("solidFill"))
        if sf is not None:
            return color_in(sf, theme, cmap)
    st = shape._element.find(f".//{_q('style', P)}")
    if st is not None:
        return color_in(st.find(_q("fillRef")), theme, cmap)
    return None


def _visible(shape):
    """Skip empty, invisible text boxes; everything else the template draws counts."""
    el = shape._element
    tag = etree.QName(el).localname
    if tag in ("pic", "grpSp", "graphicFrame", "cxnSp"):
        return True
    has_text = bool(shape.has_text_frame and shape.text_frame.text.strip())
    sppr = el.find(_q("spPr", P))
    no_fill = sppr is not None and sppr.find(_q("noFill")) is not None
    no_line = sppr is not None and sppr.find(f"{_q('ln')}/{_q('noFill')}") is not None
    styled = el.find(_q("style", P)) is not None
    filled = sppr is not None and (sppr.find(_q("solidFill")) is not None or sppr.find(_q("gradFill")) is not None
                                   or sppr.find(_q("blipFill")) is not None)
    return has_text or filled or (styled and not no_fill) or (sppr is not None and sppr.find(_q("ln")) is not None and not no_line)


def graphics(prs, layout, theme):
    """Non-placeholder shapes that show on a page using this layout: [{x, y, w, h, kind, color}]."""
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    cmap = clr_map(prs.slide_master._element, layout._element)
    src = []
    if layout._element.get("showMasterSp", "1") not in ("0", "false"):
        src += list(prs.slide_master.shapes)
    src += list(layout.shapes)
    out = []
    for sh in src:
        if sh.is_placeholder or sh.width is None or sh.height is None:
            continue
        if not _visible(sh):
            continue
        x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        if w * h > 0.6 * W * H:
            continue  # a full-page picture or panel is the background, not something to avoid
        kind = etree.QName(sh._element).localname
        out.append({"x": round(x, 3), "y": round(y, 3), "w": round(w, 3), "h": round(h, 3), "kind": kind,
                    "name": sh.name, "color": _fill_of(sh, theme, cmap) if kind == "sp" else None})
    return out


def free_area(gfx, W, H):
    """The rectangle the template's own graphics leave free. Bars and logos along the top push the
    content down, footers push it up, side bands push it in; something in the middle is reported."""
    x0, y0, x1, y1 = 0.0, 0.0, W, H
    ignored = []
    for g in gfx:
        gx0, gy0, gx1, gy1 = g["x"], g["y"], g["x"] + g["w"], g["y"] + g["h"]
        cx, cy = (gx0 + gx1) / 2, (gy0 + gy1) / 2
        if cy < 0.3 * H and g["h"] < 0.35 * H:
            y0 = max(y0, gy1 + GAP)
        elif cy > 0.7 * H and g["h"] < 0.35 * H:
            y1 = min(y1, gy0 - GAP)
        elif cx < 0.2 * W and g["w"] < 0.25 * W:
            x0 = max(x0, gx1 + GAP)
        elif cx > 0.8 * W and g["w"] < 0.25 * W:
            x1 = min(x1, gx0 - GAP)
        else:
            ignored.append(g["name"])
    return [round(x0, 3), round(y0, 3), round(x1, 3), round(y1, 3)], ignored


# ------------------------------------------------------------------ choosing layouts
def _layout_type(layout):
    return layout._element.get("type", "")


def pick_layouts(prs):
    ls = list(prs.slide_layouts)

    def by_name(kind):
        for i, l in enumerate(ls):
            n = (l.name or "").lower()
            if any(k in n for k in NAMES[kind]):
                return i
        return None

    def by_type(*types):
        for t in types:
            for i, l in enumerate(ls):
                if _layout_type(l) == t:
                    return i
        return None

    content = by_type("blank")
    if content is None:
        content = by_name("content")
    if content is None:
        content = by_type("titleOnly", "obj")
    if content is None:
        content = min(range(len(ls)), key=lambda i: len(ls[i].placeholders))
    cover = by_type("title")
    if cover is None:
        cover = by_name("cover")
    section = by_type("secHead")
    if section is None:
        section = by_name("section")
    closing = by_name("closing")
    return {"content": content, "cover": cover, "section": section, "closing": closing}


# ------------------------------------------------------------------ rendering (optional)
def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def render_layouts(prs_path, idxs, out_dir, dpi=60):
    """One PNG per layout with nothing on it but the template's own graphics (needs LibreOffice)."""
    exe = soffice()
    if not exe:
        return {}
    prs = open_template(prs_path)
    drop_slides(prs)
    order = []
    for i in idxs:
        s = prs.slides.add_slide(prs.slide_layouts[i])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        order.append(i)
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "layouts.pptx"
        prs.save(str(f))
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, str(f)],
                           capture_output=True, timeout=180)
            subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(Path(td) / "layouts.pdf"), str(Path(td) / "p")],
                           capture_output=True, timeout=120)
        except (subprocess.SubprocessError, FileNotFoundError):
            return {}
        pngs = sorted(Path(td).glob("p-*.png"))
        out = {}
        out_dir.mkdir(parents=True, exist_ok=True)
        for i, p in zip(order, pngs):
            dst = out_dir / f"layout_{i}.png"
            shutil.copy(p, dst)
            out[i] = str(dst)
        return out


def sample_color(png, box, W, H):
    """Median colour inside the free area of a rendered layout."""
    from PIL import Image
    im = Image.open(png).convert("RGB")
    sx, sy = im.width / W, im.height / H
    x0, y0, x1, y1 = box
    crop = im.crop((int(x0 * sx), int(y0 * sy), max(int(x1 * sx), int(x0 * sx) + 1), max(int(y1 * sy), int(y0 * sy) + 1)))
    crop = crop.resize((40, 24))
    px = sorted(list(getattr(crop, "get_flattened_data", crop.getdata)()), key=lambda c: sum(c))
    r, g, b = px[len(px) // 2]
    return f"{r:02X}{g:02X}{b:02X}"


# ------------------------------------------------------------------ palette
def _sat(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)[2]


def _hue(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)[0] * 360


def chromatic(h):
    return _sat(h) > 0.25 and 0.08 < rel_lum(h) < 0.85


def brand_colors(theme, gfx_all, bgs):
    """Primary and accent. Colours the template actually paints (bars, shapes, coloured cover
    backgrounds) beat theme accents, which are often PowerPoint's defaults nobody changed."""
    weight = {}
    for g in gfx_all:
        if g.get("color") and chromatic(g["color"]):
            weight[g["color"]] = weight.get(g["color"], 0) + g["w"] * g["h"] + 1
    for c in bgs:
        if c and chromatic(c):
            weight[c] = weight.get(c, 0) + 30
    accents = [theme[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if theme.get(k)]
    painted = sorted(weight, key=weight.get, reverse=True)
    if painted and accents and accents[0].upper() in OFFICE_ACCENTS:
        accents = []  # an unbranded Office theme: its blue and orange would clash with the painted brand colour
    primary = painted[0] if painted else (accents[0] if accents else "1F5FBF")
    pool = painted[1:] + accents
    accent = next((c for c in pool if chromatic(c) and min(abs(_hue(c) - _hue(primary)), 360 - abs(_hue(c) - _hue(primary))) > 35),
                  None) or next((c for c in pool if c != primary), mix(primary, "000000", 0.35))
    return primary, accent, bool(painted)


def derive_palette(page, theme, primary, accent, cmap):
    """A full palette for one kind of page in the template."""
    dark = rel_lum(page) < 0.25
    light_ink = theme.get("lt1", "FFFFFF")
    dark_ink = theme.get("dk1", "000000")
    if dark:
        text = light_ink if contrast(light_ink, page) >= 7 else "FFFFFF"
    else:
        tx = theme.get(cmap.get("tx1", "dk1"), dark_ink)
        text = tx if contrast(tx, page) >= 7 else ("111111" if contrast("111111", page) >= 7 else dark_ink)
    p = primary
    if contrast(p, page) < 2.2:
        # the brand colour is the page colour (a red cover): marks and numbers go in the ink colour
        p = text if contrast(accent, page) < 3 else accent
    pdark = mix(p, "000000", 0.3) if not dark else mix(p, "FFFFFF", 0.25)
    on_primary = "FFFFFF" if contrast("FFFFFF", p) >= 4.5 else text
    if dark:
        surface, tint, line = mix(page, "FFFFFF", 0.08), mix(p, page, 0.75), mix(text, page, 0.75)
    else:
        surface = "FFFFFF" if rel_lum(page) < 0.97 else mix(p, "FFFFFF", 0.96)
        tint, line = mix(p, "FFFFFF", 0.9), mix(text, page, 0.85)
    closing = pdark if contrast("FFFFFF", pdark) >= 4.5 else mix(pdark, "000000", 0.4)
    return {"primary": p, "primary_dark": pdark, "accent": accent, "bg_top": page, "bg_bottom": page,
            "surface": surface, "tint": tint, "text": text, "muted": mix(text, page, 0.32),
            "on_primary": on_primary, "line": line, "closing": closing,
            "on_dark": "FFFFFF" if contrast("FFFFFF", closing) >= 4.5 else text}


# ------------------------------------------------------------------ read / apply
def cmd_read(a):
    work = Path(a.work)
    tdir = work / "template"
    tdir.mkdir(parents=True, exist_ok=True)
    prs = open_template(a.template)
    drop_slides(prs)
    dst = tdir / "template.pptx"
    prs.save(str(dst))
    prs = Presentation(str(dst))
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    theme, fonts = theme_of(prs)
    picks = pick_layouts(prs)
    used = sorted({i for i in picks.values() if i is not None})
    pngs = render_layouts(dst, used, tdir) if not a.no_render else {}
    layouts, notes, gfx_all, bgs = {}, [], [], []
    for i in used:
        lay = prs.slide_layouts[i]
        gfx = graphics(prs, lay, theme)
        gfx_all += gfx
        box, ignored = free_area(gfx, W, H)
        bg, how = layout_bg(prs, lay, theme)
        if i in pngs:
            sampled = sample_color(pngs[i], box, W, H)
            if bg is None or how == "image":
                bg, how = sampled, "rendered"
        if bg is None:
            bg, how = "FFFFFF", "guess"
            notes.append(f"版式「{lay.name}」的背景是图片，没装 LibreOffice 看不到颜色，先按白底处理，请在 template.json 里改 bg")
        bgs.append(bg)
        for n in ignored:
            notes.append(f"版式「{lay.name}」页面中间有图形「{n}」，内容可能压在它上面，确认一下")
        layouts[str(i)] = {"name": lay.name, "type": _layout_type(lay), "bg": bg, "bg_from": how, "free": box,
                           "graphics": gfx, "preview": pngs.get(i)}
    primary, accent, painted = brand_colors(theme, gfx_all, bgs[1:] if len(bgs) > 1 else [])
    if not painted and theme.get("accent1", "").upper() in OFFICE_ACCENTS:
        notes.append("模板的主题色还是 PowerPoint 默认的蓝橙色，母版上也没有彩色图形，主色可能不是公司色，请在 template.json 里改 primary")
    if abs(W / H - 16 / 9) > 0.02:
        notes.append(f"模板是 {W:.2f} x {H:.2f} 英寸（不是 16:9），内容会按比例缩小放进可用区域，上下会留白")
    tj = {"source": str(Path(a.template).resolve()), "file": str(dst.resolve()), "size": [round(W, 3), round(H, 3)],
          "theme": theme, "fonts": fonts, "primary": primary, "accent": accent,
          "use": {k: (str(v) if v is not None else None) for k, v in picks.items()},
          "layouts": layouts, "notes": notes}
    save_json(tj, work / "template.json")
    print(f"wrote {work / 'template.json'}  {len(prs.slide_layouts)} 个版式，内容页用「{prs.slide_layouts[picks['content']].name}」")
    for n in notes:
        print("NOTE", n)
    return cmd_apply(a)


def build_style(work, tj, base=None):
    work = Path(work)
    prs = Presentation(tj["file"])
    cmap_master = clr_map(prs.slide_master._element)
    theme = tj["theme"]
    use = tj["use"]
    content = tj["layouts"][use["content"]]
    page = content["bg"]
    pal = derive_palette(page, theme, tj["primary"], tj["accent"], cmap_master)
    fonts_t = tj.get("fonts", {})

    def face(*keys, default):
        for k in keys:
            f = (fonts_t.get(k) or "").strip()
            if f and not f.startswith("+"):
                return f
        return default
    fonts = {"heading_ea": face("majorFont_ea", "minorFont_ea", default="Microsoft YaHei"),
             "body_ea": face("minorFont_ea", "majorFont_ea", default="Microsoft YaHei"),
             "heading_latin": face("majorFont_latin", default="Arial"),
             "body_latin": face("minorFont_latin", default="Arial")}
    fonts["display_latin"] = fonts["heading_latin"]
    style = load_json(base) if base else {}
    style.update({
        "name": "template", "label": f"公司模板 {Path(tj['source']).stem}",
        "palette": pal, "fonts": fonts,
        "shape": {"radius": 0.06, "shadow": False, "card_alpha": 1.0, "depth": "whisper"},
        # the template's master is the running header; a centred series name on top would only
        # collide with its logo and push every page's content down
        "decor": {"header": False, "leaf": False, "brush": False, "sparkles": False, "title_mark": True,
                  "closing_dark": True, "bullet": "icon", "emphasis": "color", "chrome": "standard"},
        "background": {"decor": [], "cover_decor": [], "closing_decor": [], "flat": True},
        "illustration": {"style_prompt": "clean flat corporate illustration, restrained, matching the brand colours",
                         "negative": "text, letters, words, logo, watermark", "fallback_look": "flat"},
        "emphasis_color": "accent" if contrast(pal["accent"], page) >= 3 else "primary",
        "rules": ["版式和配色跟着公司模板走，母版上的 logo、色条、页脚不要遮挡"], "forbid": {},
    })
    # each kind of page sits on its own template layout; a layout whose background differs from the
    # content pages (a red cover, a dark closing) gets its own palette
    kinds = {}
    for kind in ("content", "cover", "section", "closing"):
        idx = use.get(kind)
        if kind == "closing" and idx is None:
            idx = use.get("cover")  # most templates end on their cover layout
        if idx is None:
            idx = use["content"]
        lay = tj["layouts"][idx]
        lay_pal = derive_palette(lay["bg"], theme, tj["primary"], tj["accent"], cmap_master)
        entry = {"index": int(idx), "free": lay["free"]}
        if kind == "closing":
            # the closing page uses its layout's background as the "closing" colour
            dark = rel_lum(lay["bg"]) < 0.25 or (contrast("FFFFFF", lay["bg"]) >= 3 and chromatic(lay["bg"]))
            lay_pal.update({"closing": lay["bg"], "on_dark": "FFFFFF" if contrast("FFFFFF", lay["bg"]) >= 4.5 else lay_pal["text"]})
            entry["closing_dark"] = dark
            entry["palette"] = lay_pal
        elif lay["bg"] != page:
            entry["palette"] = lay_pal
        kinds[kind] = entry
    style["template"] = {"file": tj["file"], "size": tj["size"], "kinds": kinds}
    return style


def cmd_apply(a):
    work = Path(a.work)
    tj = load_json(work / "template.json")
    st = build_style(work, tj, getattr(a, "base", None))
    save_json(st, work / "style.json")
    p = st["palette"]
    print(f"wrote {work / 'style.json'}  主色 #{p['primary']} 强调 #{p['accent']} 正文 #{p['text']} 底色 #{p['bg_top']}")
    return cmd_card(a)


# ------------------------------------------------------------------ card
def cmd_card(a):
    from PIL import Image, ImageDraw
    from common import get_font as cjk_font
    work = Path(a.work)
    tj = load_json(work / "template.json")
    st = load_json(work / "style.json")
    W, H = tj["size"]
    Wc = 1600
    f_t, f_b, f_s = cjk_font(34, True), cjk_font(22), cjk_font(18)
    tw = 360
    th = int(tw * H / W)
    kinds = st["template"]["kinds"]
    rows = 2 + (len(tj["notes"]) + 1) // 2
    img = Image.new("RGB", (Wc, 140 + th + 120 + 220 + 40 * rows), "#F4F5F7")
    d = ImageDraw.Draw(img)
    d.text((48, 40), "公司模板确认卡", font=f_t, fill="#1F2937")
    d.text((48, 92), f"{Path(tj['source']).name}   {W:.2f} x {H:.2f} 英寸", font=f_s, fill="#5B6474")
    x = 48
    labels = {"cover": "封面", "section": "章节页", "content": "内容页", "closing": "结尾页"}
    for kind in ("cover", "section", "content", "closing"):
        k = kinds[kind]
        lay = tj["layouts"][str(k["index"])]
        y = 140
        if lay.get("preview") and Path(lay["preview"]).exists():
            im = Image.open(lay["preview"]).convert("RGB").resize((tw, th))
        else:
            im = Image.new("RGB", (tw, th), "#" + lay["bg"])
        img.paste(im, (x, y))
        fx0, fy0, fx1, fy1 = k["free"]
        sx, sy = tw / W, th / H
        d.rectangle((x + fx0 * sx, y + fy0 * sy, x + fx1 * sx, y + fy1 * sy), outline="#E11D48", width=2)
        d.rectangle((x, y, x + tw, y + th), outline="#C9CED6", width=1)
        d.text((x, y + th + 10), f"{labels[kind]}  ·  {lay['name']}", font=f_s, fill="#1F2937")
        x += tw + 24
    y = 140 + th + 60
    d.text((48, y), "红框是内容能放的区域，母版上的 logo、色条和页脚都留在框外", font=f_s, fill="#5B6474")
    y += 50
    p = st["palette"]
    sw = [("主色", p["primary"]), ("强调色", p["accent"]), ("正文", p["text"]), ("底色", p["bg_top"]),
          ("卡片", p["surface"]), ("结尾页", kinds["closing"].get("palette", p).get("closing", p["closing"]))]
    x = 48
    for name, c in sw:
        d.rounded_rectangle((x, y, x + 120, y + 80), 10, fill="#" + c, outline="#C9CED6")
        d.text((x, y + 90), f"{name} #{c}", font=f_s, fill="#1F2937")
        x += 150
    f = st["fonts"]
    d.text((x + 20, y + 4), f"标题字体 {f['heading_ea']} / {f['heading_latin']}", font=f_b, fill="#1F2937")
    d.text((x + 20, y + 44), f"正文字体 {f['body_ea']} / {f['body_latin']}", font=f_b, fill="#1F2937")
    y += 150
    for n in tj["notes"]:
        d.text((48, y), "· " + n, font=f_s, fill="#B45309")
        y += 34
    out = work / "template" / "template_card.png"
    img = img.crop((0, 0, Wc, min(img.height, y + 30)))
    img.save(out)
    print(f"wrote {out}")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("read")
    r.add_argument("template")
    r.add_argument("work")
    r.add_argument("--base", help="a style.json whose layout choices to keep")
    r.add_argument("--no-render", action="store_true", help="skip the LibreOffice previews")
    p = sub.add_parser("apply")
    p.add_argument("work")
    p.add_argument("--base")
    c = sub.add_parser("card")
    c.add_argument("work")
    a = ap.parse_args()
    return {"read": cmd_read, "apply": cmd_apply, "card": cmd_card}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main() or 0)
