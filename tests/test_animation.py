"""Slide transitions and item builds.

Run with: python -m unittest discover -s tests -v
"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

import animate as A
import layouts as L


def run_script(name, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          env=dict(os.environ, PYTHONIOENCODING="utf-8"))


def blank_slide(n_boxes=3):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ids = [slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(3), Inches(0.8)).shape_id
           for i in range(n_boxes)]
    return prs, slide, ids


def child_tags(slide):
    return [c.tag.split("}")[1] for c in slide._element]


class ResolveTests(unittest.TestCase):
    def test_defaults_and_overrides(self):
        self.assertEqual(A.resolve({}, {}), ("transition", "rise"))
        self.assertEqual(A.resolve({"animation": "click"}, {}), ("click", "rise"))
        self.assertEqual(A.resolve({"animation": "click"}, {"animation": "none"}), ("none", "rise"))
        self.assertEqual(A.resolve({"animation": {"mode": "auto", "effect": "fade"}}, {}), ("auto", "fade"))
        self.assertEqual(A.resolve({"animation": True}, {}), ("auto", "rise"))
        self.assertEqual(A.resolve({"animation": "sparkle"}, {}), ("transition", "rise"))
        self.assertEqual(A.resolve({"animation_effect": "spin"}, {}), ("transition", "rise"))

    def test_scenario_picks_the_default(self):
        self.assertEqual(A.resolve({"scenario": "oral"}, {})[0], "click")
        self.assertEqual(A.resolve({"scenario": "talk"}, {})[0], "click")
        self.assertEqual(A.resolve({"scenario": "report"}, {})[0], "transition")
        self.assertEqual(A.resolve({"scenario": "oral", "animation": "auto"}, {})[0], "auto")


class WriterTests(unittest.TestCase):
    def test_click_builds_one_group_per_click(self):
        prs, slide, ids = blank_slide(3)
        A.set_transition(slide)
        self.assertEqual(A.set_builds(slide, [[ids[0]], [ids[1], ids[2]]], "click"), 2)
        timing = slide._element.find(qn("p:timing"))
        main = next(timing.iter(qn("p:seq"))).find(qn("p:cTn")).find(qn("p:childTnLst"))
        self.assertEqual(len(main.findall(qn("p:par"))), 2)
        types = [c.get("nodeType") for c in timing.iter(qn("p:cTn")) if c.get("presetClass") == "entr"]
        self.assertEqual(types, ["clickEffect", "clickEffect", "withEffect"])
        self.assertEqual(A.animated_ids(slide), set(ids))
        self.assertEqual(child_tags(slide)[-2:], ["transition", "timing"])
        ctn_ids = [int(c.get("id")) for c in timing.iter(qn("p:cTn"))]
        self.assertEqual(len(ctn_ids), len(set(ctn_ids)))

    def test_auto_starts_with_page_and_staggers(self):
        prs, slide, ids = blank_slide(3)
        A.set_builds(slide, [[i] for i in ids], "auto", "fade")
        timing = slide._element.find(qn("p:timing"))
        self.assertTrue(any(c.get("evt") == "onBegin" for c in timing.iter(qn("p:cond"))))
        effects = [c for c in timing.iter(qn("p:cTn")) if c.get("presetClass") == "entr"]
        self.assertEqual([e.get("nodeType") for e in effects], ["afterEffect", "withEffect", "withEffect"])
        delays = [int(e.find(qn("p:stCondLst"))[0].get("delay")) for e in effects]
        self.assertEqual(delays, [0, A.STAGGER, 2 * A.STAGGER])

    def test_single_group_or_still_modes_write_no_timing(self):
        for mode, groups in (("click", [[1]]), ("transition", [[1], [2]]), ("none", [[1], [2]])):
            with self.subTest(mode=mode):
                prs, slide, ids = blank_slide(2)
                self.assertEqual(A.set_builds(slide, groups, mode), 0)
                self.assertIsNone(slide._element.find(qn("p:timing")))

    def test_none_removes_transition(self):
        prs, slide, ids = blank_slide(1)
        A.set_transition(slide)
        A.set_transition(slide, on=False)
        self.assertIsNone(slide._element.find(qn("p:transition")))

    def test_apply_skips_untagged_and_chrome(self):
        prs, slide, ids = blank_slide(4)
        rendered = [({"t": "text"}, [ids[0]]), ({"t": "rect", "step": 0}, [ids[1]]),
                    ({"t": "text", "step": 1}, [ids[2]]), ({"t": "text", "step": 1, "role": "chrome"}, [ids[3]])]
        self.assertEqual(A.apply(slide, rendered, "click"), 2)
        self.assertEqual(A.animated_ids(slide), {ids[1], ids[2]})


class LayoutTagTests(unittest.TestCase):
    def test_item_loop_elements_carry_their_step(self):
        L.step_reset()
        L.step(0)
        a = L.T("x", 0, 0, 1, 1, 12)
        L.step()
        b = L.R(0, 0, 1, 1)
        L.step_end()
        c = L.T("y", 0, 0, 1, 1, 12)
        self.assertEqual((a.get("step"), b.get("step"), c.get("step")), (0, 1, None))


class EndToEndTests(unittest.TestCase):
    def build(self, deck):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, True)
        shutil.copy(ROOT / "references" / "styles" / "business-blue.json", tmp / "style.json")
        (tmp / "deck.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
        for args in (("plan_deck.py", tmp), ("build_deck.py", tmp)):
            r = run_script(*args)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        dump = json.loads((tmp / "out" / "layout_dump.json").read_text(encoding="utf-8"))
        return Presentation(str(tmp / "out" / "beautified.pptx")), dump

    def test_click_deck_builds_items_not_titles(self):
        items = [{"title": f"t{i}", "text": f"text {i}"} for i in range(3)]
        prs, dump = self.build({"meta": {"animation": "click"}, "slides": [
            {"n": 1, "type": "cover", "title": "cover"},
            {"n": 2, "type": "process", "title": "steps", "items": items},
            {"n": 3, "type": "list", "title": "list", "items": items, "animation": "none"}]})
        self.assertEqual([s["animation"]["builds"] for s in dump["slides"]], [0, 3, 0])
        cover, steps, still = prs.slides
        self.assertIsNotNone(cover._element.find(qn("p:transition")))
        self.assertIsNone(still._element.find(qn("p:transition")))
        animated = A.animated_ids(steps)
        self.assertTrue(animated <= {sh.shape_id for sh in steps.shapes})
        title = next(sh for sh in steps.shapes if sh.has_text_frame and sh.text_frame.text == "steps")
        self.assertNotIn(title.shape_id, animated)

    def test_default_is_transition_only(self):
        prs, dump = self.build({"meta": {}, "slides": [
            {"n": 1, "type": "list", "title": "list", "items": [{"text": "a"}, {"text": "b"}]}]})
        self.assertEqual(dump["slides"][0]["animation"], {"mode": "transition", "builds": 0})
        self.assertIsNotNone(prs.slides[0]._element.find(qn("p:transition")))


if __name__ == "__main__":
    unittest.main()
