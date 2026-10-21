"""Regression checks for extraction, fidelity normalization and pipeline status.

Run with: python -m unittest discover -s tests -v
"""
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from build_deck import Builder
from common import wrap_lines
from extract_deck import para_text_marked
import pipeline
from qa_check import norm


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def run_script(name, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          env=dict(os.environ, PYTHONIOENCODING="utf-8"))


class FidelityTests(unittest.TestCase):
    def test_preserves_numeric_signs(self):
        for text in ("-5°C", "-12.5%", "–5°C", "−5°C"):
            with self.subTest(text=text):
                self.assertEqual(norm(text), text)
        self.assertEqual(norm("first\n-5°C"), "first-5°C")
        self.assertEqual(norm("-\n5°C"), "-5°C")

    def test_still_ignores_list_markers_and_bold(self):
        for marker in ("·", "•", "●", "▪", "◦", "*", "-", "–"):
            with self.subTest(marker=marker):
                self.assertEqual(norm(f"{marker} **annual report**"), "annualreport")
        self.assertEqual(norm("- first\n- second"), "firstsecond")
        self.assertEqual(norm("•-5°C"), "-5°C")

    def test_reports_negative_value_changed_to_positive(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            out = work / "out"
            out.mkdir()
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            write_json(work / "style.json", {"name": "test"})
            write_json(work / "deck.json", {"meta": {}, "slides": [{"n": 1, "type": "content"}]})
            write_json(work / "source_text.json", {"slides": [
                {"n": 1, "units": [{"id": "s1_1_p0", "text": "-5°C"}]}]})
            write_json(out / "layout_dump.json", {"slides": [{"n": 1, "elements": []}]})
            for actual, expected_code in (("5°C", 1), ("-5°C", 0)):
                with self.subTest(actual=actual):
                    box.text = actual
                    prs.save(str(out / "beautified.pptx"))
                    result = run_script("qa_check.py", work)
                    self.assertEqual(result.returncode, expected_code, result.stdout + result.stderr)
                    report = json.loads((out / "qa_report.json").read_text(encoding="utf-8"))
                    errors = [i for i in report["issues"] if i["level"] == "ERROR"]
                    self.assertEqual(len(errors), expected_code)
                    if errors:
                        self.assertEqual(errors[0]["check"], "fidelity")


class ExtractionTests(unittest.TestCase):
    def test_soft_breaks_preserve_partial_and_full_bold(self):
        for left_bold, right_bold, expected in (
                (False, False, "annual\vreport"),
                (True, True, "annual\vreport"),
                (True, False, "**annual**\vreport"),
                (False, True, "annual\v**report**")):
            with self.subTest(left_bold=left_bold, right_bold=right_bold):
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                para = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame.paragraphs[0]
                left = para.add_run()
                left.text, left.font.bold = "annual", left_bold
                para.add_line_break()
                para.add_run()  # an empty run must not shift the text/bold mapping
                right = para.add_run()
                right.text, right.font.bold = "report", right_bold
                self.assertEqual(para_text_marked(para), expected)

    def test_plain_paragraph_and_break_only(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        para = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame.paragraphs[0]
        para.text = "annual report"
        self.assertEqual(para_text_marked(para), "annual report")
        para.clear()
        para.add_line_break()
        self.assertEqual(para_text_marked(para), "\v")

    def test_extracts_soft_break_to_source_and_draft(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            para = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame.paragraphs[0]
            para.add_run().text = "annual"
            para.add_line_break()
            para.add_run().text = "report"
            src = work / "input.pptx"
            prs.save(str(src))
            result = run_script("extract_deck.py", src, work / "extracted", "--no-render")
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            for filename, key in (("source_text.json", "units"), ("deck.draft.json", "units"), ("raw.json", "shapes")):
                with self.subTest(filename=filename):
                    data = json.loads((work / "extracted" / filename).read_text(encoding="utf-8"))
                    self.assertEqual(data["slides"][0][key][0]["text"], "annual\vreport")


class SoftBreakRoundTripTests(unittest.TestCase):
    def _para(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        return slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame.paragraphs[0]

    def test_builder_turns_vertical_tab_into_soft_break(self):
        para = self._para()
        runs = Builder._add_runs(para, "annual\vreport")
        self.assertEqual([r.text for r in runs], ["annual", "report"])
        tags = [child.tag.split("}")[1] for child in para._p if child.tag.split("}")[1] in ("r", "br")]
        self.assertEqual(tags, ["r", "br", "r"])
        self.assertEqual(para_text_marked(para), "annual\vreport")

    def test_soft_break_inherits_run_size(self):
        para = self._para()
        for r in Builder._add_runs(para, "annual\vreport"):
            r.font.size = Pt(28)
        Builder._sync_breaks(para)
        br = para._p.find(qn("a:br"))
        self.assertEqual(br.find(qn("a:rPr")).get("sz"), "2800")

    def test_paragraph_break_stays_separate_from_soft_break(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
        Builder._add_runs(tf.paragraphs[0], "one\vtwo")
        Builder._add_runs(tf.add_paragraph(), "three")
        self.assertEqual("\n".join(para_text_marked(p) for p in tf.paragraphs), "one\vtwo\nthree")
        self.assertEqual(norm("one\vtwo\nthree"), "onetwothree")

    def test_wrap_counts_soft_break_as_new_line(self):
        self.assertEqual(len(wrap_lines("a\vb", 20, 500)), 2)


class PipelineTests(unittest.TestCase):
    def test_keeps_review_and_returns_qa_status(self):
        for qa_status in (0, 1, 2):
            with self.subTest(qa_status=qa_status), tempfile.TemporaryDirectory() as tmp:
                called = []

                def run_stage(args):
                    called.append(Path(args[0]).name)
                    if called[-1] != "qa_check.py":
                        return 0
                    if qa_status != 2:  # a finished QA run leaves its report; a crash does not
                        (Path(tmp) / "out").mkdir(exist_ok=True)
                        (Path(tmp) / "out" / "qa_report.json").write_text("{}", encoding="utf-8")
                    return qa_status

                with patch.object(sys, "argv", ["pipeline.py", tmp, "--lite"]), \
                        patch.object(pipeline, "run", side_effect=run_stage), \
                        contextlib.redirect_stdout(io.StringIO()):
                    status = pipeline.main()
                self.assertEqual(status, qa_status)
                self.assertEqual(called[-1], "review.py")


if __name__ == "__main__":
    unittest.main()
