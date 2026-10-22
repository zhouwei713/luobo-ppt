"""Company templates: decks built on a customer's own .pptx / .potx, master and theme kept.

Run with: python -m unittest discover -s tests -v
"""
import copy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from common import contrast

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU = 914400
RED = "C8102E"

DECK = {"meta": {"title": "测试", "footer": "测试部"}, "slides": [
    {"n": 1, "type": "cover", "title": "模板测试稿", "subtitle": "副标题"},
    {"n": 2, "type": "list", "title": "三件事", "items": [{"title": "一", "text": "第一件"},
                                                       {"title": "二", "text": "第二件"},
                                                       {"title": "三", "text": "第三件"}]},
    {"n": 3, "type": "chart", "title": "增长", "chart": {"type": "column", "categories": ["Q1", "Q2", "Q3"],
                                                        "series": [{"name": "A", "values": [1, 2, 3]}]},
     "source": "示例"},
    {"n": 4, "type": "closing", "headline": "谢谢", "body": "再见"}]}


def run_script(name, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          env=dict(os.environ, PYTHONIOENCODING="utf-8"))


def set_theme(prs, colors, ea=None):
    tp = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(tp.blob)
    cs = root.find(f".//{{{A}}}clrScheme")
    for k, v in colors.items():
        el = cs.find(f"{{{A}}}{k}")
        for c in list(el):
            el.remove(c)
        etree.SubElement(el, f"{{{A}}}srgbClr", val=v)
    if ea:
        for f in root.iter(f"{{{A}}}majorFont", f"{{{A}}}minorFont"):
            f.find(f"{{{A}}}ea").set("typeface", ea)
    tp._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def to_master(prs, target, draw):
    """Draw on a scratch slide, then move the shapes into the master or a layout."""
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    draw(tmp)
    for sp in list(tmp.shapes._spTree)[2:]:
        target.shapes._spTree.append(copy.deepcopy(sp))
    lst = prs.slides._sldIdLst
    sid = lst[-1]
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)


def bar(s, x, y, w, h, color):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(color)
    r.line.fill.background()


def label(s, text, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(10)


def make_template(path, dark=False, potx=False):
    """A company template: a red bar and a logo along the top, a footer line at the bottom, a red
    cover layout, Office's default theme accents (so the red has to come from the master), and two
    sample slides that must not end up in the deck."""
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    set_theme(p, {"accent1": "4472C4", "accent2": "ED7D31"}, "微软雅黑")
    if dark:
        p.slide_master.background.fill.solid()
        p.slide_master.background.fill.fore_color.rgb = RGBColor.from_string("12163A")
        cm = p.slide_master._element.find(qn("p:clrMap"))
        cm.set("bg1", "dk1")
        cm.set("tx1", "lt1")

    def draw(s):
        bar(s, 0, 0, 13.333, 0.12, RED)
        bar(s, 11.0, 0.3, 1.8, 0.5, RED)  # the logo block
        bar(s, 0.5, 6.95, 12.333, 0.02, RED)
        label(s, "内部资料", 0.5, 7.0, 4, 0.3)
    to_master(p, p.slide_master, draw)
    if not dark:
        tl = p.slide_layouts[0]
        tl.background.fill.solid()
        tl.background.fill.fore_color.rgb = RGBColor.from_string(RED)
    for i in (0, 1):
        s = p.slides.add_slide(p.slide_layouts[i])
        s.shapes.title.text = "示例页，请删除"
    p.save(str(path))
    if potx:
        data = Path(path).read_bytes()
        out = Path(path).with_suffix(".potx")
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                b = zin.read(it.filename)
                if it.filename == "[Content_Types].xml":
                    b = b.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
                zout.writestr(it, b)
        return out
    return path


class TemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, True)

    def build(self, name, dark=False, potx=False):
        t = make_template(self.tmp / f"{name}.pptx", dark=dark, potx=potx)
        w = self.tmp / name
        r = run_script("template_kit.py", "read", t, w, "--no-render")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        (w / "deck.json").write_text(json.dumps(DECK, ensure_ascii=False), encoding="utf-8")
        for args in (("plan_deck.py", w, "--level", "full"), ("make_assets.py", w), ("build_deck.py", w)):
            r = run_script(*args)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        r = run_script("qa_check.py", w)
        rep = json.loads((w / "out" / "qa_report.json").read_text(encoding="utf-8"))
        return w, rep["issues"]

    def test_template_reading(self):
        t = make_template(self.tmp / "read.pptx")
        w = self.tmp / "read"
        r = run_script("template_kit.py", "read", t, w, "--no-render")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        tj = json.loads((w / "template.json").read_text(encoding="utf-8"))
        # the red painted on the master beats the untouched Office accents
        self.assertEqual(tj["primary"], RED)
        self.assertNotIn(tj["accent"], ("4472C4", "ED7D31"))
        content = tj["layouts"][tj["use"]["content"]]
        x0, y0, x1, y1 = content["free"]
        self.assertGreater(y0, 0.8)    # below the logo block (0.3 + 0.5)
        self.assertLess(y1, 6.95)      # above the footer line
        cover = tj["layouts"][tj["use"]["cover"]]
        self.assertEqual(cover["bg"], RED)
        st = json.loads((w / "style.json").read_text(encoding="utf-8"))
        self.assertEqual(st["fonts"]["heading_ea"], "微软雅黑")
        self.assertIn("palette", st["template"]["kinds"]["cover"])  # the red cover has its own palette
        self.assertTrue((w / "template" / "template_card.png").exists())

    def test_deck_on_template(self):
        w, issues = self.build("light")
        self.assertEqual([i for i in issues if i["level"] == "ERROR"], [])
        self.assertEqual([i for i in issues if i["check"] == "contrast"], [])
        prs = Presentation(str(w / "out" / "beautified.pptx"))
        self.assertEqual(len(prs.slides), 4)  # the template's sample slides are gone
        names = [s.slide_layout.name for s in prs.slides]
        self.assertEqual(names[0], "Title Slide")
        self.assertEqual(names[1], "Blank")
        # the master's own graphics are still there, untouched
        self.assertEqual(len([s for s in prs.slide_master.shapes if not s.is_placeholder]), 4)
        tj = json.loads((w / "template.json").read_text(encoding="utf-8"))
        x0, y0, x1, y1 = tj["layouts"][tj["use"]["content"]]["free"]
        for s in list(prs.slides)[1:3]:
            self.assertEqual(len(s.placeholders), 0)  # no empty "click to add title" boxes
            for sh in s.shapes:
                self.assertGreaterEqual(sh.top / EMU, y0 - 0.02, sh.name)
                self.assertLessEqual((sh.top + sh.height) / EMU, y1 + 0.02, sh.name)
                self.assertGreaterEqual(sh.left / EMU, x0 - 0.02, sh.name)
                self.assertLessEqual((sh.left + sh.width) / EMU, x1 + 0.02, sh.name)
        # the cover sits on the red layout: its text is light
        dump = json.loads((w / "out" / "layout_dump.json").read_text(encoding="utf-8"))
        cover = dump["slides"][0]
        self.assertEqual(cover["template"]["kind"], "cover")
        title = next(e for e in cover["elements"] if e["t"] == "text" and e["text"] == "模板测试稿")
        ink = cover["template"]["palette"][title["color"]] if title["color"] in cover["template"]["palette"] else None
        if ink:
            self.assertGreaterEqual(contrast(ink, RED), 3.0)

    def test_dark_template(self):
        w, issues = self.build("dark", dark=True)
        self.assertEqual([i for i in issues if i["level"] == "ERROR"], [])
        self.assertEqual([i for i in issues if i["check"] == "contrast"], [])
        st = json.loads((w / "style.json").read_text(encoding="utf-8"))
        self.assertGreaterEqual(contrast(st["palette"]["text"], "12163A"), 7)

    def test_potx(self):
        w, issues = self.build("potx", potx=True)
        self.assertEqual([i for i in issues if i["level"] == "ERROR"], [])
        self.assertTrue((w / "out" / "beautified.pptx").exists())

    def test_lite_preview_follows_the_template(self):
        w, _ = self.build("lite")
        r = run_script("render_lite.py", w)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        from PIL import Image
        im = Image.open(w / "out" / "png" / "slide-01.png").convert("RGB")
        r_, g_, b_ = im.getpixel((im.width // 2, im.height - 12))
        self.assertGreater(r_, 150)  # the red cover background shows through
        self.assertLess(g_, 80)


if __name__ == "__main__":
    unittest.main()
